"""Bygger VIRKE_Protokoll fra styremøte (generalforsamling).docx + .dotx med v3 stilark.

Viktig endring fra forrige iterasjon:
  Meta-blokka (Dato/Sted/Tid/Til stede/Forfall/Referent) er nå *flytetekst med
  Fremhevet-labels*, ikke en tabell — iht. Figma-mønsteret i A4-eksempelet.

Kjøres med:
    cd /Users/afg/Nettsteder/virke-word
    python3 scripts/virke-build/build_protokoll.py
"""
import sys, os
HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

from docx import Document
from virke_stylesheet import (
    install_virke_stylesheet, setup_page, add_header_footer,
    add_tittel, add_h1, add_h2, add_h3, add_body, add_helptext,
    add_placeholder_to, add_body_with_placeholder, add_bold_to,
    style_virke_table, save_as_dotx,
)

OUT_DIR = "/Users/afg/Nettsteder/virke-word/word-maler/Eksempler"
OUT_DOCX = os.path.join(OUT_DIR, "VIRKE_Protokoll fra styremøte (generalforsamling).docx")
OUT_DOTX = os.path.join(OUT_DIR, "VIRKE_Protokoll fra styremøte (generalforsamling).dotx")


def _meta_line(doc, label, placeholder):
    """En linje i meta-blokka: 'Label: [placeholder]' med label i bold."""
    p = doc.add_paragraph(style="Virke Brødtekst")
    add_bold_to(p, f"{label}: ")
    add_placeholder_to(p, placeholder)
    return p


def build_document():
    doc = Document()
    install_virke_stylesheet(doc)
    setup_page(doc)
    add_header_footer(doc, section_label="Selskapsrett",
                      doc_title="Protokoll fra styremøte")

    # ── Tittel med inline plassholder ──
    p = doc.add_paragraph(style="Virke Tittel")
    p.add_run("Protokoll fra styremøte i ")
    add_placeholder_to(p, "[foretaksnavn]")

    # ── Meta-blokk (flytetekst, ikke tabell) ──
    _meta_line(doc, "Dato", "[dd.mm.åååå]")
    _meta_line(doc, "Sted", "[Sted eller digital plattform]")

    # Tid har to plassholdere
    p = doc.add_paragraph(style="Virke Brødtekst")
    add_bold_to(p, "Tid: ")
    add_bold_to(p, "Møtestart ")
    add_placeholder_to(p, "[tt:mm]")
    p.add_run("   ")
    add_bold_to(p, "Møteslutt ")
    add_placeholder_to(p, "[tt:mm]")

    _meta_line(doc, "Til stede", "[Navn og ev. stilling på tilstedeværende]")
    _meta_line(doc, "Forfall", "[Navn på innkalte som hadde forfall]")
    _meta_line(doc, "Referent", "[Navn på referenten]")

    # ── Konstituering ──
    add_body(doc,
        "Styret var beslutningsdyktig. Styreleder ledet møtet. Det var ingen "
        "innvendinger til innkallingen eller dagsorden. Møtet ble erklært for "
        "lovlig satt.")
    add_body(doc, "Følgende saker forelå til behandling:")

    # ── Sak 1 — Årsregnskap ──
    add_h2(doc, "Sak 1 — Årsregnskap")
    add_body_with_placeholder(doc, [
        ("text", "Det ble fremlagt forslag til årsregnskap for regnskapsåret "),
        ("placeholder", "[årstall]"),
        ("text", "."),
    ])
    add_body(doc,
        "Forslag til årsberetning, resultatregnskap, balanse og noter ble "
        "gjennomgått. Styret besluttet enstemmig å foreslå følgende disponering "
        "av årets overskudd:")

    # Disponering-tabell
    disp = doc.add_table(rows=4, cols=2)
    disp.rows[0].cells[0].text = ""; disp.rows[0].cells[1].text = ""
    disp.rows[0].cells[0].paragraphs[0].add_run("Post")
    disp.rows[0].cells[1].paragraphs[0].add_run("Beløp")

    rows_data = [
        ("Utbytte",           "kr [1 000 000],-"),
        ("Annen egenkapital", "kr [500 000],-"),
        ("Sum disponert",     "kr [1 500 000],-"),
    ]
    for i, (label, val) in enumerate(rows_data, start=1):
        l = disp.rows[i].cells[0]; r = disp.rows[i].cells[1]
        l.text = ""; r.text = ""
        lr = l.paragraphs[0].add_run(label)
        rr = r.paragraphs[0].add_run(val)
        if label == "Sum disponert":
            lr.bold = True; rr.bold = True
    style_virke_table(disp, header=True)

    add_body(doc,
        "Styret vedtok enstemmig forslaget til årsregnskap, den fremlagte "
        "årsberetningen, resultatregnskapets disponeringer, balanse og noter "
        "som selskapets regnskap.")
    add_helptext(doc,
        "Kort tekst om året kan legges inn her. Eksempel: stabil omsetning, "
        "nye løsninger, nyansettelser, utsikter for kommende år. Slett eller "
        "tilpass denne boksen.")
    add_body(doc,
        "Styret vedtok å fremlegge årsregnskapet m.v. for generalforsamlingen "
        "for godkjenning.")

    # ── Sak 2 — Honorar ──
    add_h2(doc, "Sak 2 — Honorar til styret")
    add_body_with_placeholder(doc, [
        ("text", "Styret besluttet å foreslå for generalforsamlingen at "
                 "styrehonoraret settes til "),
        ("placeholder", "[beløp/sats]"),
        ("text", "."),
    ])

    # ── Sak 3 — Innkalling ──
    add_h2(doc, "Sak 3 — Innkalling til generalforsamling")
    add_body_with_placeholder(doc, [
        ("text", "Styret besluttet å innkalle til ordinær generalforsamling i "
                 "selskapet for behandling av angitte saker. Generalforsamlingen "
                 "innkalles til "),
        ("placeholder", "[dato]"),
        ("text", " og avholdes "),
        ("placeholder", "[møtested / digital plattform]"),
        ("text", "."),
    ])

    # ── Avslutning ──
    add_body(doc, "Alle beslutninger var enstemmige.")
    add_body(doc, "Det var ingen flere saker til behandling. Møtet ble hevet.")
    add_body_with_placeholder(doc, [
        ("text", "Neste styremøte er i henhold til styreplan: "),
        ("placeholder", "[dato]"),
        ("text", "."),
    ])

    # ── Signaturblokk ──
    add_h1(doc, "Signaturer")
    sp = doc.add_paragraph(style="Virke Brødtekst")
    add_placeholder_to(sp, "[Sted / dato]")

    sig = doc.add_table(rows=3, cols=3)
    for i in range(3):
        c = sig.rows[0].cells[i]; c.text = ""
        c.paragraphs[0].add_run("_______________________________")
    for i in range(3):
        c = sig.rows[1].cells[i]; c.text = ""
        add_placeholder_to(c.paragraphs[0], "[Fornavn Etternavn]")
    for i, role in enumerate(["Styreleder", "Styremedlem", "Styremedlem"]):
        c = sig.rows[2].cells[i]; c.text = ""
        c.paragraphs[0].add_run(role)
    style_virke_table(sig, header=False, borders=False)

    return doc


def main():
    doc = build_document()
    doc.save(OUT_DOCX)
    print(f"SAVED DOCX: {OUT_DOCX}")

    doc2 = build_document()
    save_as_dotx(doc2, OUT_DOTX)
    print(f"SAVED DOTX: {OUT_DOTX}")


if __name__ == "__main__":
    main()
