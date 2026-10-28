"""Bygger VIRKE_Aksjeeierbok.docx + .dotx med v3 stilark.

Kjøres med:
    cd /Users/afg/Nettsteder/virke-word
    python3 scripts/virke-build/build_aksje.py
"""
import sys, os
HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from virke_stylesheet import (
    install_virke_stylesheet, setup_page, add_header_footer,
    add_tittel, add_h1, add_h2, add_h3, add_body, add_helptext,
    add_placeholder_to, add_body_with_placeholder,
    style_virke_table, save_as_dotx,
)

OUT_DIR = "/Users/afg/Nettsteder/virke-word/word-maler/Eksempler"
OUT_DOCX = os.path.join(OUT_DIR, "VIRKE_Aksjeeierbok.docx")
OUT_DOTX = os.path.join(OUT_DIR, "VIRKE_Aksjeeierbok.dotx")


def build_document():
    doc = Document()
    install_virke_stylesheet(doc)
    setup_page(doc)
    add_header_footer(doc, section_label="Selskapsrett",
                      doc_title="Aksjeeierbok")

    # ── Tittel ──
    add_tittel(doc, "Aksjeeierbok")

    # ── Intro ──
    add_body(doc,
        "Når selskapet er stiftet, skal styret uten opphold opprette en aksjeeierbok. "
        "I denne skal aksjeeierne innføres i alfabetisk orden med angivelse av navn, "
        "fødselsnummer eller organisasjonsnummer, adresse, antall aksjer og nummer på "
        "aksjene. Det må også oppgis hvilken klasse aksjene tilhører dersom "
        "aksjeselskapet har flere aksjeklasser.")

    add_body(doc,
        "Når en aksjeeier er innført i aksjeeierboken, skal selskapet gi aksjeeieren "
        "melding om dette. Den skal også være tilgjengelig for enhver som vil ha "
        "innsyn. Aksjeeierboken oppdateres ved eierskifte eller pantsettelse av "
        "aksjer etter fastsatte regler i aksjeloven (§ 4-7). Selskapet må oppbevare "
        "opplysninger om tidligere aksjeeiere i minst ti år.")

    add_helptext(doc,
        "Dette dokumentet er av generell art. Teksten må tilpasses den enkelte "
        "virksomhet. Slett denne boksen før dokumentet tas i bruk.")

    # ── Aksjeeierregister ──
    add_h1(doc, "Aksjeeierregister")

    headers = [
        "Navn",
        "Fødselsnr. / org.nr.",
        "Adresse",
        "Antall aksjer",
        "Aksjenr. (fra – til)",
        "Aksjeklasse",
        "Dato innføring",
    ]
    table = doc.add_table(rows=1, cols=len(headers))
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        cell.paragraphs[0].add_run(h)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    placeholder_row = [
        "[Navn]",
        "[Fødselsnr./org.nr.]",
        "[Adresse]",
        "[Antall]",
        "[fra–til]",
        "[Klasse]",
        "[dd.mm.åååå]",
    ]
    for _ in range(5):  # 5 blanke rader for utfylling
        row = table.add_row()
        for i, cell in enumerate(row.cells):
            cell.text = ""
            add_placeholder_to(cell.paragraphs[0], placeholder_row[i])

    style_virke_table(table, header=True, zebra=True)

    # ── Overdragelser og merknader ──
    add_h1(doc, "Overdragelser og merknader")

    add_body(doc,
        "Hvert eierskifte eller pantsettelse skal innføres i aksjeeierboken "
        "uten ugrunnet opphold etter at selskapet har fått melding om det.")

    merk = doc.add_table(rows=4, cols=3)
    merk_headers = ["Dato", "Type hendelse", "Kommentar"]
    for i, h in enumerate(merk_headers):
        cell = merk.rows[0].cells[i]
        cell.text = ""
        cell.paragraphs[0].add_run(h)

    for rix in range(1, 4):
        for cix, ph in enumerate(["[dd.mm.åååå]", "[Eierskifte / pantsettelse / ...]", "[fritekst]"]):
            cell = merk.rows[rix].cells[cix]
            cell.text = ""
            add_placeholder_to(cell.paragraphs[0], ph)

    style_virke_table(merk, header=True, zebra=True)

    # ── Signaturblokk ──
    add_h1(doc, "Bekreftelse")
    add_body_with_placeholder(doc, [
        ("text", "Denne aksjeeierboken bekreftes som korrekt pr. "),
        ("placeholder", "[dd.mm.åååå]"),
        ("text", "."),
    ])

    sig = doc.add_table(rows=3, cols=2)
    for i in range(2):
        c = sig.rows[0].cells[i]; c.text = ""
        c.paragraphs[0].add_run("_______________________________")
    for i, n in enumerate(["[Fornavn Etternavn]", "[Fornavn Etternavn]"]):
        c = sig.rows[1].cells[i]; c.text = ""
        add_placeholder_to(c.paragraphs[0], n)
    for i, role in enumerate(["Styreleder", "Styremedlem"]):
        c = sig.rows[2].cells[i]; c.text = ""
        c.paragraphs[0].add_run(role)
    style_virke_table(sig, header=False, borders=False)

    return doc


def main():
    doc = build_document()
    doc.save(OUT_DOCX)
    print(f"SAVED DOCX: {OUT_DOCX}")

    # For .dotx bygger vi dokumentet på nytt, slik at vi ikke muterer det
    # allerede lagrede (python-docx holder referanse til xml-treet).
    doc2 = build_document()
    save_as_dotx(doc2, OUT_DOTX)
    print(f"SAVED DOTX: {OUT_DOTX}")


if __name__ == "__main__":
    main()
